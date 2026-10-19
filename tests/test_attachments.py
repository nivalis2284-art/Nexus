"""Real extraction, original preservation, owner isolation and paged tool reads."""

import io
import json
from pathlib import Path

import pytest

from nexus.attachments import AttachmentStore, MAX_FILE_BYTES, message_content
from nexus.tools import build


def save(store, name, data, owner="s1"):
    return store.save("sessions", owner, name, "", io.BytesIO(data), len(data))


@pytest.mark.parametrize("name,raw", [
    ("notes.md", "# 中文说明\nkeep me\nnext".encode()),
    ("script.py", b"print('hello')\n"),
    ("data.csv", "姓名,数量\n甲,3".encode("gb18030")),
    ("wide.txt", "文本".encode("utf-16")),
])
def test_text_original_and_extracted_content_survive_reopen(tmp_path, name, raw):
    store = AttachmentStore(tmp_path)
    item = save(store, name, raw)
    reopened = AttachmentStore(tmp_path)
    restored = reopened.get("sessions", "s1", item["id"])
    assert Path(restored["original_path"]).read_bytes() == raw
    assert restored["status"] == "ready"
    assert reopened.read("sessions", "s1", item["id"], 1, 100)["content"]


def test_duplicate_names_paths_and_owner_isolation(tmp_path):
    store = AttachmentStore(tmp_path)
    first = save(store, "../../same.txt", b"first")
    second = save(store, "same.txt", b"second")
    assert first["name"] == "same.txt"
    assert first["original_path"] != second["original_path"]
    with pytest.raises(ValueError):
        store.resolve("sessions", "other", [first["id"]])
    with pytest.raises(ValueError):
        store.get("sessions", "s1", "../s2")
    with pytest.raises(ValueError):
        store.resolve("sessions", "s1", [first["id"], first["id"]])


def test_document_extractors_include_structural_content(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    word = Document()
    word.add_paragraph("Word paragraph")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Word table"
    word.save(tmp_path / "word.docx")
    book = Workbook()
    book.active.title = "统计"
    book.active.append(["sales", 10, "=B1*2"])
    book.save(tmp_path / "book.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Slide text"
    slide.shapes.add_table(1, 1, 0, Inches(1), Inches(4), Inches(1)).table.cell(0, 0).text = "Slide table"
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    slides.save(tmp_path / "slides.pptx")
    expected = {
        "word.docx": ["Word paragraph", "Word table"],
        "book.xlsx": ["Sheet: 统计", "A1=sales", "C1==B1*2"],
        "slides.pptx": ["Slide 1", "Slide text", "Slide table", "Speaker notes"],
    }
    store = AttachmentStore(tmp_path / "home")
    for name, pieces in expected.items():
        item = save(store, name, (tmp_path / name).read_bytes())
        assert item["status"] == "ready", item
        text = store.read("sessions", "s1", item["id"], 1, 100)["content"]
        assert all(piece in text for piece in pieces), text


def pdf_bytes(*, encrypted=False, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (PDF attachment) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_text_scanned_and_password_cases_are_distinct(tmp_path):
    store = AttachmentStore(tmp_path)
    text = save(store, "text.pdf", pdf_bytes())
    assert "[Page 1]" in store.read("sessions", "s1", text["id"], 1, 100)["content"]
    assert "PDF attachment" in store.read("sessions", "s1", text["id"], 1, 100)["content"]
    scanned = save(store, "scan.pdf", pdf_bytes(blank=True))
    assert scanned["status"] == "empty"
    assert "OCR" in scanned["note"]
    password = save(store, "protected.pdf", pdf_bytes(encrypted=True))
    assert password["status"] == "error"
    assert "密码" in password["note"]
    assert Path(password["original_path"]).exists()


def test_unsupported_corrupt_oversize_and_interrupted_upload(tmp_path):
    store = AttachmentStore(tmp_path)
    unsupported = save(store, "archive.zip", b"PK\x00\x01")
    assert unsupported["status"] == "stored"
    damaged = save(store, "bad.docx", b"not a document")
    assert damaged["status"] == "error"
    assert Path(damaged["original_path"]).read_bytes() == b"not a document"
    registry, context = build(root=tmp_path, attachment_source=(store, "sessions", "s1"),
                              attachment_ids={unsupported["id"], damaged["id"]})
    for item in (unsupported, damaged):
        result = registry.execute("read_attachment", {"attachment_id": item["id"]}, context)
        assert not result.ok
        assert item["original_path"] in result.content
    before = list(store.owner_path("sessions", "s1").iterdir())
    with pytest.raises(ValueError):
        store.save("sessions", "s1", "large.bin", "", io.BytesIO(), MAX_FILE_BYTES + 1)
    with pytest.raises(ValueError, match="未完成"):
        store.save("sessions", "s1", "partial.txt", "", io.BytesIO(b"a"), 100)
    assert list(store.owner_path("sessions", "s1").iterdir()) == before


def test_tool_can_read_sent_attachment_in_pages_but_not_unsent_file(tmp_path):
    store = AttachmentStore(tmp_path)
    item = save(store, "long.txt", b"first\nsecond\nthird")
    unsent = save(store, "unsent.txt", b"do not reveal")
    registry, context = build(root=tmp_path, attachment_source=(store, "sessions", "s1"),
                              attachment_ids={item["id"]})
    first = registry.execute("read_attachment", {"attachment_id": item["id"], "limit": 1}, context)
    assert first.ok and "1: first" in first.content and "Next offset: 2" in first.content
    second = registry.execute("read_attachment", {"attachment_id": item["id"], "offset": 2}, context)
    assert second.ok and "2: second" in second.content
    assert not registry.execute("read_attachment", {"attachment_id": unsent["id"]}, context).ok


def test_long_single_line_is_readable_across_pages(tmp_path):
    store = AttachmentStore(tmp_path)
    item = save(store, "minified.js", b"a" * 20000 + b"tail marker")
    result = store.read("sessions", "s1", item["id"], 1, 1)
    assert result["next_offset"] == 2
    assert "tail marker" in store.read("sessions", "s1", item["id"], 2, 1)["content"]


def test_task_image_content_only_when_model_supports_it(tmp_path):
    store = AttachmentStore(tmp_path)
    item = store.save("tasks", "task_1", "pixel.png", "image/png", io.BytesIO(b"png"), 3)
    assert isinstance(message_content("image", [item]), str)
    content = message_content("image", [item], include_images=True)
    assert content[1]["type"] == "image"
    assert item["id"] in json.dumps(content)


def test_extraction_limit_is_reported_and_original_is_complete(tmp_path, monkeypatch):
    import nexus.attachments as module
    monkeypatch.setattr(module, "MAX_TEXT_CHARS", 16)
    store = AttachmentStore(tmp_path)
    data = b"large document body that remains available"
    item = save(store, "large.txt", data)
    assert item["truncated"] is True and "上限" in item["note"]
    assert Path(item["original_path"]).read_bytes() == data
    text = store.read("sessions", "s1", item["id"], 1, 100)["content"]
    assert "large document b" in text
