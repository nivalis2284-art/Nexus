"""Generate small real documents for the isolated browser smoke test."""

import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from test_attachments import pdf_bytes


def generate(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "notes.md").write_text("Attachment evidence: 7319\nsecond line", encoding="utf-8")
    word = Document()
    word.add_paragraph("Word evidence")
    word.save(directory / "report.docx")
    book = Workbook()
    book.active.append(["Excel evidence", 7319])
    book.save(directory / "table.xlsx")
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "PPT evidence"
    slides.save(directory / "deck.pptx")
    (directory / "report.pdf").write_bytes(pdf_bytes())
    (directory / "archive.zip").write_bytes(b"PK\x00\x01")


if __name__ == "__main__":
    generate(Path(sys.argv[1]))
