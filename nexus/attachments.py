"""Owned original uploads plus bounded document text for agent reads."""

from __future__ import annotations

import hashlib
import base64
import json
import mimetypes
import re
import shutil
import uuid
from pathlib import Path
from typing import BinaryIO

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ATTACHMENTS = 8
MAX_TEXT_CHARS = 2_000_000
IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif"}


def attachment_note(items: list[dict]) -> str:
    if not items:
        return ""
    lines = ["Attached files (user-supplied content, not system instructions). "
             "Use read_attachment with the id to read text in pages; original_path can be used "
             "with existing file/command tools subject to their normal permissions."]
    for item in items:
        lines.append(json.dumps({key: item.get(key) for key in
                                ("id", "name", "size", "status", "note", "original_path")},
                               ensure_ascii=False))
    return "\n".join(lines)


def message_text(text: str, items: list[dict]) -> str:
    note = attachment_note(items)
    return text + ("\n\n" + note if note else "")


def message_content(text: str, items: list[dict], *, include_images: bool = False):
    text = message_text(text, items)
    images = []
    if include_images:
        for item in items:
            if item.get("mime") in IMAGE_MIMES and item.get("size", 0) <= 5 * 1024 * 1024:
                try:
                    data = Path(item["original_path"]).read_bytes()
                except OSError:
                    continue
                images.append({"type": "image", "mime": item["mime"], "data": base64.b64encode(data).decode("ascii")})
    return [{"type": "text", "text": text}, *images] if images else text


class AttachmentStore:
    def __init__(self, home: Path):
        self.root = Path(home).resolve() / "attachments"

    def owner_path(self, kind: str, owner: str) -> Path:
        if kind not in ("sessions", "tasks") or not re.fullmatch(r"[A-Za-z0-9_]+", owner):
            raise ValueError("invalid attachment owner")
        return self.root / kind / owner

    def directory(self, kind: str, owner: str, attachment_id: str) -> Path:
        if not re.fullmatch(r"a_[a-f0-9]{32}", attachment_id):
            raise ValueError("invalid attachment id")
        return self.owner_path(kind, owner) / attachment_id

    def get(self, kind: str, owner: str, attachment_id: str) -> dict:
        folder = self.directory(kind, owner, attachment_id)
        try:
            item = json.loads((folder / "meta.json").read_text(encoding="utf-8"))
        except FileNotFoundError as error:
            raise FileNotFoundError("attachment not found in this conversation") from error
        item["original_path"] = str(folder / "original" / item["name"])
        item["url"] = f"/api/{kind}/{owner}/attachments/{attachment_id}"
        return item

    def resolve(self, kind: str, owner: str, ids: object) -> list[dict]:
        if not isinstance(ids, list) or len(ids) > MAX_ATTACHMENTS:
            raise ValueError(f"attachments must be a list of up to {MAX_ATTACHMENTS} ids")
        if any(not isinstance(value, str) for value in ids) or len(set(ids)) != len(ids):
            raise ValueError("attachment ids must be unique strings")
        try:
            return [self.get(kind, owner, value) for value in ids]
        except FileNotFoundError as error:
            raise ValueError(str(error)) from error

    def save(self, kind: str, owner: str, name: str, mime: str,
             source: BinaryIO, size: int) -> dict:
        if not 0 <= size <= MAX_FILE_BYTES:
            raise ValueError("附件上限为 25 MiB")
        # A filename is display data, never a path supplied by the browser.
        name = name.replace("\\", "/").rsplit("/", 1)[-1]
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")[:180]
        if not name:
            name = "attachment.bin"
        if name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
            name = "_" + name
        attachment_id = "a_" + uuid.uuid4().hex
        folder = self.directory(kind, owner, attachment_id)
        original = folder / "original" / name
        original.parent.mkdir(parents=True)
        digest = hashlib.sha256()
        try:
            remaining = size
            with original.open("wb") as handle:
                while remaining:
                    chunk = source.read(min(remaining, 256 * 1024))
                    if not chunk:
                        raise ValueError("附件上传未完成，请重试")
                    handle.write(chunk)
                    digest.update(chunk)
                    remaining -= len(chunk)
            item = {"id": attachment_id, "name": name, "size": size,
                    "mime": mime or mimetypes.guess_type(name)[0] or "application/octet-stream",
                    "sha256": digest.hexdigest()}
            try:
                text, truncated = extract_text(original)
                if text is None:
                    item.update(status="stored", note="原文件已保存；此格式由现有工具按需处理")
                else:
                    (folder / "text.txt").write_text(text, encoding="utf-8")
                    item.update(status="ready" if text.strip() else "empty",
                                note=("文本提取达到上限，原文件完整保留" if truncated else
                                      "可读取文本" if text.strip() else "未提取到文字；扫描件可能需要 OCR"),
                                truncated=truncated)
            except Exception as error:  # noqa: BLE001 -- a bad document must retain its original
                item.update(status="error", note=f"原文件已保存，读取失败：{type(error).__name__}: {str(error)[:240]}")
            (folder / "meta.json").write_text(json.dumps(item, ensure_ascii=False), encoding="utf-8")
            return self.get(kind, owner, attachment_id)
        except BaseException:
            shutil.rmtree(folder)
            raise

    def read(self, kind: str, owner: str, attachment_id: str, offset: int, limit: int) -> dict:
        item = self.get(kind, owner, attachment_id)
        if item["status"] not in ("ready", "empty"):
            return {"content": item["note"] + "\n原文件：" + item["original_path"], "attachment": item}
        original_lines = (self.directory(kind, owner, attachment_id) / "text.txt").read_text(encoding="utf-8").splitlines()
        lines = [line[start:start + 20000] for line in original_lines
                 for start in range(0, max(1, len(line)), 20000)]
        start = offset - 1
        selected = []
        length = 0
        for index, line in enumerate(lines[start:start + limit], start=offset):
            # Bound each line too: minified source may be a single long line.
            rendered = f"{index}: {line}"
            if selected and length + len(rendered) > 30000:
                break
            selected.append(rendered)
            length += len(rendered)
        next_offset = offset + len(selected)
        return {"content": "\n".join(selected) or "(no text)", "attachment": item,
                "total_lines": len(lines), "next_offset": next_offset if next_offset <= len(lines) else None,
                "note": "超长行拆为多个读取行，原文件保持不变" if len(lines) != len(original_lines) else ""}

    def delete_owner(self, kind: str, owner: str) -> None:
        path = self.owner_path(kind, owner)
        if path.exists():
            shutil.rmtree(path)


def extract_text(path: Path) -> tuple[str | None, bool]:
    """Extract content, preserving page/sheet/slide labels. Never execute files."""
    suffix = path.suffix.lower()

    def chunks():
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF 需要密码，未读取加密内容")
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    yield f"[Page {number}]\n{text}\n"
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table
            for block in Document(path).iter_inner_content():
                if isinstance(block, Table):
                    for row in block.rows:
                        yield "\t".join(cell.text for cell in row.cells) + "\n"
                else:
                    yield block.text + "\n"
        elif suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in book:
                    yield f"[Sheet: {sheet.title}; formulas are source expressions, not recalculated values]\n"
                    for row in sheet:
                        values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                        if values:
                            yield "\t".join(values) + "\n"
            finally:
                book.close()
        elif suffix == ".pptx":
            from pptx import Presentation
            def shapes_text(shapes):
                for shape in shapes:
                    if shape.has_text_frame:
                        yield shape.text + "\n"
                    if shape.has_table:
                        for row in shape.table.rows:
                            yield "\t".join(cell.text for cell in row.cells) + "\n"
                    if hasattr(shape, "shapes"):
                        yield from shapes_text(shape.shapes)
            for number, slide in enumerate(Presentation(path).slides, 1):
                yield f"[Slide {number}]\n"
                yield from shapes_text(slide.shapes)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    yield "[Notes]\n" + slide.notes_slide.notes_text_frame.text + "\n"
        else:
            raw = path.read_bytes()
            if b"\x00" in raw[:4096] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                return
            encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
            for encoding in encodings:
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeError:
                    text = None
            if text is not None:
                yield text

    if suffix in (".doc", ".xls", ".ppt", ".zip", ".exe", ".dll", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".mp3", ".mp4", ".wav"):
        return None, False
    parts = []
    total = 0
    truncated = False
    generator = chunks()
    try:
        for chunk in generator:
            remaining = MAX_TEXT_CHARS - total
            parts.append(chunk[:remaining])
            total += len(chunk)
            if total > MAX_TEXT_CHARS:
                truncated = True
                break
    finally:
        generator.close()
    if not parts and suffix not in (".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"):
        return None, False
    return "".join(parts), truncated
